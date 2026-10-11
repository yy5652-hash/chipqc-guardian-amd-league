"""Defensive corpus walk and per-type parsers.

Every parser is wrapped per file: an unreadable, encrypted, unknown or corrupt
file is recorded and skipped, never allowed to stop the walk. Text parsers run
in worker threads with a timeout; anything that needs the vision model (images,
image-only PDF pages) is queued and done afterwards on the GPU.
"""
from __future__ import annotations

import ast
import csv
import io
import logging
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, TimeoutError as FutTimeout
from dataclasses import dataclass, field
from pathlib import Path

from . import config

log = logging.getLogger("mc3.ingest")

TEXT_EXT = {
    ".txt", ".log", ".md", ".markdown", ".rst", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".ini", ".cfg", ".conf", ".toml", ".properties", ".env.example", ".sh", ".bash", ".zsh", ".bat", ".ps1",
    ".js", ".ts", ".tsx", ".jsx", ".c", ".h", ".cpp", ".hpp", ".cc", ".java", ".kt", ".go", ".rs", ".rb",
    ".sql", ".tsv", ".tex", ".adoc", ".org", ".csv.txt", ".out", ".err", ".trace",
}
CODE_EXT = {".py", ".pyi"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx", ".docm"}
XLSX_EXT = {".xlsx", ".xlsm"}
CSV_EXT = {".csv"}

STATUS_WORDS = r"(WITHDRAWN|SUPERSEDED|OBSOLETE|DEPRECATED|RETIRED|ARCHIVED|REVOKED|CANCELLED|CANCELED|RESCINDED|VOID|OLD|DRAFT|PRELIMINARY)"
STATUS_FILENAME_RE = re.compile(r"(?:^|[_\-\s.\[(])" + STATUS_WORDS + r"(?:$|[_\-\s.\])])", re.I)
STATUS_TEXT_RE = re.compile(r"\b" + STATUS_WORDS + r"\b")  # case-sensitive: headline markers are upper case


@dataclass
class Segment:
    text: str
    loc: str = ""
    kind: str = "text"            # text | table | row | code | image | image_page
    image_path: str | None = None  # for image kinds: path to open at query time
    page: int | None = None


@dataclass
class FileDoc:
    rel: str
    abs: str
    ext: str
    status: str = ""
    segments: list[Segment] = field(default_factory=list)
    error: str = ""
    kind: str = "text"
    ocr_pages: list[int] = field(default_factory=list)  # pdf pages needing OCR (0-based)


# --------------------------------------------------------------------------- helpers

def read_text_bytes(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("latin-1", errors="replace")


def read_text_file(p: Path) -> str:
    with open(p, "rb") as fh:
        data = fh.read(config.MAX_TEXT_CHARS * 2)
    return read_text_bytes(data)[: config.MAX_TEXT_CHARS]


def detect_status(name: str, head_text: str = "") -> str:
    m = STATUS_FILENAME_RE.search(name)
    if m:
        return m.group(1).upper()
    head = (head_text or "")[:600]
    m = STATUS_TEXT_RE.search(head)
    if m:
        w = m.group(1).upper()
        # "DRAFT"/"OLD" are too common inside prose; only trust them in the filename
        if w not in {"DRAFT", "OLD", "PRELIMINARY"}:
            return w
    return ""


def line_segments(text: str, lines_per: int, overlap: int, kind: str, loc_prefix: str = "lines") -> list[Segment]:
    lines = text.splitlines()
    if not lines:
        return []
    segs: list[Segment] = []
    i = 0
    n = len(lines)
    while i < n:
        j = min(n, i + lines_per)
        block = "\n".join(lines[i:j]).strip()
        if block:
            segs.append(Segment(text=block, loc=f"{loc_prefix} {i + 1}-{j}", kind=kind))
        if j >= n:
            break
        i = max(i + 1, j - overlap)
    return segs


def _cells(row) -> list[str]:
    return ["" if v is None else " ".join(str(v).split()) for v in row]


def rows_to_segments(rows: list[list[str]], loc_prefix: str, header: list[str] | None) -> list[Segment]:
    """Turn a table into row-level lines grouped into small segments.

    Each line keeps the header association ("col: value; col: value") so that a
    lexical hit on an identifier lands on the row that carries its value.
    """
    segs: list[Segment] = []
    if not rows:
        return segs
    lines: list[tuple[int, str]] = []
    for ri, r in enumerate(rows, start=1):
        if header and len(header) >= 2:
            pairs = []
            for h, v in zip(header, r):
                if v != "":
                    pairs.append(f"{h}: {v}" if h else v)
            extra = [v for v in r[len(header):] if v != ""]
            pairs.extend(extra)
            if pairs:
                lines.append((ri, f"[{loc_prefix} row {ri}] " + "; ".join(pairs)))
        else:
            vals = [v for v in r if v != ""]
            if vals:
                lines.append((ri, f"[{loc_prefix} row {ri}] " + " | ".join(vals)))
    if not lines:
        return segs
    hdr = ""
    if header and len(header) >= 2:
        hdr = f"[{loc_prefix} columns] " + " | ".join(h for h in header if h) + "\n"
    per = config.ROWS_PER_CHUNK
    for i in range(0, len(lines), per):
        block = lines[i:i + per]
        text = hdr + "\n".join(t for _, t in block)
        segs.append(Segment(text=text, loc=f"{loc_prefix} rows {block[0][0]}-{block[-1][0]}", kind="row"))
    return segs


def pick_header(rows: list[list[str]]) -> tuple[list[str] | None, list[list[str]]]:
    """First row with >=2 non-empty cells is the header, if it looks textual."""
    for idx, r in enumerate(rows[:5]):
        nonempty = [c for c in r if c != ""]
        if len(nonempty) >= 2:
            numeric = sum(1 for c in nonempty if re.fullmatch(r"[-+]?\d[\d,]*(\.\d+)?", c))
            if numeric <= len(nonempty) // 2:
                return r, rows[idx + 1:]
            return None, rows
    return None, rows


# --------------------------------------------------------------------------- parsers

def parse_pdf(p: Path, doc: FileDoc) -> None:
    import fitz  # PyMuPDF

    d = fitz.open(str(p))
    try:
        if d.is_encrypted:
            if d.needs_pass and not d.authenticate(""):
                doc.error = "encrypted (password required)"
                return
            if d.needs_pass:
                doc.error = "encrypted (password required)"
                return
        n = min(d.page_count, config.MAX_PDF_PAGES)
        first_text = ""
        for i in range(n):
            try:
                page = d[i]
                txt = page.get_text("text") or ""
            except Exception as e:  # a single bad page must not kill the file
                log.warning("pdf %s page %d: %s", doc.rel, i + 1, e)
                txt = ""
            if i == 0:
                first_text = txt
            # tables, when PyMuPDF can see them, as explicit rows
            table_text = ""
            try:
                tabs = page.find_tables()
                for t in tabs.tables[:6]:
                    rows = [_cells(r) for r in t.extract()]
                    rows = [r for r in rows if any(r)]
                    if len(rows) >= 2:
                        table_text += "\n".join(" | ".join(r) for r in rows) + "\n"
            except Exception:
                pass
            if len(txt.strip()) >= 25:
                doc.segments.append(Segment(text=txt.strip(), loc=f"page {i + 1}", kind="text", page=i + 1))
                if table_text.strip():
                    doc.segments.append(Segment(text=table_text.strip(), loc=f"page {i + 1} table", kind="table", page=i + 1))
            else:
                # image-only / scanned page: OCR later with the vision model
                if len(doc.ocr_pages) < config.MAX_OCR_PAGES_PER_PDF:
                    doc.ocr_pages.append(i)
        if d.page_count > n:
            log.warning("pdf %s truncated to %d pages", doc.rel, n)
        doc.status = detect_status(p.name, first_text)
        try:
            meta = d.metadata or {}
            title = (meta.get("title") or "").strip()
            if title and not doc.status:
                doc.status = detect_status(title)
        except Exception:
            pass
    finally:
        d.close()


def parse_docx(p: Path, doc: FileDoc) -> None:
    import docx  # python-docx

    d = docx.Document(str(p))
    paras: list[str] = []
    order: list[tuple[str, object]] = []
    try:
        for item in d.iter_inner_content():
            order.append((item.__class__.__name__, item))
    except AttributeError:
        order = [("Paragraph", x) for x in d.paragraphs] + [("Table", x) for x in d.tables]
    buf: list[str] = []
    head_text = ""

    def flush():
        if buf:
            text = "\n".join(buf).strip()
            if text:
                doc.segments.extend(
                    Segment(text=s.text, loc=f"paragraphs {s.loc}", kind="text")
                    for s in line_segments(text, 40, 4, "text", "")
                )
            buf.clear()

    for kind, item in order:
        if kind == "Paragraph":
            t = " ".join(item.text.split())
            if t:
                buf.append(t)
                if len(head_text) < 600:
                    head_text += t + "\n"
        else:  # Table
            rows: list[list[str]] = []
            for r in item.rows:
                cells: list[str] = []
                prev = None
                for c in r.cells:
                    ct = " ".join(c.text.split())
                    if ct == prev:  # merged cells repeat the same cell object
                        continue
                    cells.append(ct)
                    prev = ct
                if any(cells):
                    rows.append(cells)
            if rows:
                flush()
                header, body = pick_header(rows)
                context = "\n".join(buf[-2:]) if buf else ""
                table_segs = rows_to_segments(body, f"table {len(doc.segments) + 1}", header)
                for s in table_segs:
                    s.kind = "table"
                    if context:
                        s.text = context + "\n" + s.text
                doc.segments.extend(table_segs)
    flush()
    # headers/footers occasionally carry document status
    try:
        for section in d.sections:
            for part in (section.header, section.footer):
                for para in part.paragraphs:
                    t = " ".join(para.text.split())
                    if t:
                        head_text += t + "\n"
    except Exception:
        pass
    doc.status = detect_status(p.name, head_text)


def parse_xlsx(p: Path, doc: FileDoc) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            try:
                for row in ws.iter_rows(values_only=True):
                    cells = _cells(row)
                    if any(cells):
                        rows.append(cells)
                    if len(rows) > 20000:
                        break
            except Exception as e:
                log.warning("xlsx %s sheet %s: %s", doc.rel, ws.title, e)
            if not rows:
                continue
            header, body = pick_header(rows)
            doc.segments.extend(rows_to_segments(body, f"sheet '{ws.title}'", header))
    finally:
        wb.close()
    doc.status = detect_status(p.name)


def parse_csv(p: Path, doc: FileDoc) -> None:
    text = read_text_file(p)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[str]] = []
    for r in reader:
        cells = [" ".join(c.split()) for c in r]
        if any(cells):
            rows.append(cells)
        if len(rows) > 50000:
            break
    header, body = pick_header(rows)
    doc.segments.extend(rows_to_segments(body, "csv", header))
    doc.status = detect_status(p.name)


def parse_code(p: Path, doc: FileDoc) -> None:
    src = read_text_file(p)
    doc.segments.extend(line_segments(src, 60, 6, "code"))
    lines = src.splitlines()
    consts: list[str] = []
    try:
        tree = ast.parse(src)
    except Exception:
        tree = None
    if tree is not None:
        def comment_above(lineno: int) -> str:
            out = []
            k = lineno - 2
            while k >= 0 and lines[k].strip().startswith("#"):
                out.insert(0, lines[k].strip())
                k -= 1
            return " ".join(out)

        def record(node, label: str):
            try:
                seg = " ".join(l.strip() for l in lines[node.lineno - 1:node.end_lineno])
            except Exception:
                return
            c = comment_above(node.lineno)
            consts.append(f"line {node.lineno}: {label}{seg}" + (f"    (comment above: {c})" if c else ""))

        for node in ast.walk(tree):
            if isinstance(node, (ast.Assign, ast.AnnAssign)):
                targets = node.targets if isinstance(node, ast.Assign) else [node.target]
                names = [t.id for t in targets if isinstance(t, ast.Name)]
                span = (getattr(node, "end_lineno", None) or node.lineno) - node.lineno + 1
                if names and span <= 6:
                    record(node, "")
            elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                args = node.args
                positional = args.posonlyargs + args.args
                defaults = args.defaults
                for a, dflt in zip(positional[len(positional) - len(defaults):], defaults):
                    try:
                        consts.append(f"line {node.lineno}: def {node.name}(...): default {a.arg} = {ast.unparse(dflt)}")
                    except Exception:
                        pass
                for a, dflt in zip(args.kwonlyargs, args.kw_defaults):
                    if dflt is not None:
                        try:
                            consts.append(f"line {node.lineno}: def {node.name}(...): default {a.arg} = {ast.unparse(dflt)}")
                        except Exception:
                            pass
    if consts:
        text = "Constants, assignments and parameter defaults in this source file:\n" + "\n".join(consts[:400])
        doc.segments.append(Segment(text=text, loc="constants and defaults", kind="code"))
    doc.status = detect_status(p.name)


def parse_text(p: Path, doc: FileDoc) -> None:
    text = read_text_file(p)
    if not text.strip():
        doc.error = "no text"
        return
    doc.segments.extend(line_segments(text, 45, 5, "text"))
    doc.status = detect_status(p.name, text[:600])


# --------------------------------------------------------------------------- walk

def walk_files(corpus: Path) -> list[Path]:
    out: list[Path] = []

    def onerror(e):
        log.warning("walk error: %s", e)

    for root, dirs, files in os.walk(corpus, onerror=onerror, followlinks=False):
        dirs.sort()
        files.sort()
        for f in files:
            out.append(Path(root) / f)
    return out


def ingest_file(path: Path, corpus: Path) -> FileDoc:
    try:
        rel = path.relative_to(corpus).as_posix()
    except Exception:
        rel = path.name
    doc = FileDoc(rel=rel, abs=str(path), ext=path.suffix.lower())
    t0 = time.monotonic()
    try:
        if path.is_symlink() and not path.exists():
            doc.error = "broken symlink"
            return doc
        if not path.is_file():
            doc.error = "not a regular file"
            return doc
        size = path.stat().st_size
        if size == 0:
            doc.error = "empty file"
            return doc
        if size > config.MAX_FILE_BYTES:
            doc.error = f"too large ({size} bytes)"
            return doc
        with open(path, "rb") as fh:  # permission probe: raises PermissionError for mode 000
            fh.read(16)
        ext = doc.ext
        name_lower = path.name.lower()
        if ext in PDF_EXT:
            doc.kind = "pdf"
            parse_pdf(path, doc)
        elif ext in DOCX_EXT:
            doc.kind = "docx"
            parse_docx(path, doc)
        elif ext in XLSX_EXT:
            doc.kind = "xlsx"
            parse_xlsx(path, doc)
        elif ext in CSV_EXT:
            doc.kind = "csv"
            parse_csv(path, doc)
        elif ext in CODE_EXT:
            doc.kind = "code"
            parse_code(path, doc)
        elif ext in IMAGE_EXT:
            doc.kind = "image"
            # transcription happens later on the GPU; record the intent now
            doc.segments.append(Segment(text="", loc="image", kind="image", image_path=str(path)))
            doc.status = detect_status(path.name)
        elif ext in TEXT_EXT or name_lower in {"readme", "changelog", "license", "makefile", "dockerfile", "notes"}:
            doc.kind = "text"
            parse_text(path, doc)
        else:
            # unknown extension: accept it only if it is clearly text
            with open(path, "rb") as fh:
                head = fh.read(4096)
            if head and b"\x00" not in head:
                try:
                    head.decode("utf-8")
                    doc.kind = "text"
                    parse_text(path, doc)
                except UnicodeDecodeError:
                    doc.error = "unknown binary type"
            else:
                doc.error = "unknown binary type"
    except PermissionError as e:
        doc.error = f"permission denied: {e}"
    except Exception as e:  # noqa: BLE001 - anything at all is skipped, never raised
        doc.error = f"{type(e).__name__}: {str(e)[:200]}"
    finally:
        log.info("ingested %s kind=%s segs=%d status=%s err=%s (%.1fs)", doc.rel, doc.kind, len(doc.segments),
                 doc.status or "-", doc.error or "-", time.monotonic() - t0)
    return doc


def ingest_corpus(corpus: Path, timeout_s: float | None = None, workers: int = 4) -> list[FileDoc]:
    """Parse every file (text parsers in threads with a timeout). Never raises."""
    timeout_s = timeout_s or config.PARSE_TIMEOUT_S
    corpus = Path(corpus)
    docs: list[FileDoc] = []
    try:
        files = walk_files(corpus)
    except Exception as e:
        log.error("walk failed: %s", e)
        files = []
    log.info("corpus %s: %d files", corpus, len(files))
    if not files:
        return docs
    pool = ThreadPoolExecutor(max_workers=workers)
    futures = [(f, pool.submit(ingest_file, f, corpus)) for f in files]
    for f, fut in futures:
        try:
            docs.append(fut.result(timeout=timeout_s))
        except FutTimeout:
            log.error("parse timeout: %s", f)
            try:
                rel = f.relative_to(corpus).as_posix()
            except Exception:
                rel = f.name
            docs.append(FileDoc(rel=rel, abs=str(f), ext=f.suffix.lower(), error="parse timeout"))
        except Exception as e:
            docs.append(FileDoc(rel=f.name, abs=str(f), ext=f.suffix.lower(), error=f"{type(e).__name__}: {e}"))
    pool.shutdown(wait=False, cancel_futures=True)
    return docs
